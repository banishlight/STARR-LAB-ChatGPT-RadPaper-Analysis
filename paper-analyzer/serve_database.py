#!/usr/bin/env python3
"""
STARR-LAB Radiation Paper Database Viewer
=========================================
Serves the spreadsheet made by rad_paper_analyzer_local.py as a web page:
one tab per sheet, with search, sorting, paging, and links to each paper's
PDF (and to extracted figure images). The spreadsheet itself can be
downloaded from the page.

Requirements:
    pip install openpyxl

Run:
    python serve_database.py --xlsx rad_effects_database.xlsx --pdf-dir ~/pdfs
Then open http://localhost:8000

All links in the page are relative, so it also works behind a reverse proxy
under a path prefix (e.g. Caddy forwarding /radpapers/ with the prefix stripped).

The spreadsheet is re-read automatically when the file changes, so a new
analyzer run shows up on refresh. Use --host 0.0.0.0 to share on your network.
"""

import argparse
import json
import mimetypes
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import load_workbook

PAGE_SIZE = 100

# Columns that hold the paper name, in the sheets build_spreadsheet writes
PAPER_COLUMNS = ("Paper Name", "Paper")


class Database:
    """The spreadsheet, loaded into memory and reloaded when the file changes."""

    def __init__(self, xlsx_path: Path):
        self.xlsx_path = xlsx_path
        self.mtime = None
        self.sheets = {}         # name -> {"headers": [...], "rows": [[...], ...]}
        self.paper_files = {}    # paper name -> PDF filename
        self.lock = threading.Lock()

    def refresh(self):
        mtime = self.xlsx_path.stat().st_mtime
        with self.lock:
            if mtime == self.mtime:
                return
            wb = load_workbook(self.xlsx_path, read_only=True, data_only=True)
            sheets = {}
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                if not rows:
                    continue
                headers = ["" if h is None else str(h) for h in rows[0]]
                body = [["" if v is None else v for v in r] for r in rows[1:] if any(v not in (None, "") for v in r)]
                sheets[ws.title] = {"headers": headers, "rows": body}
            wb.close()

            # Fallback for sheets without a Source File column: map paper names to files
            paper_files = {}
            papers = sheets.get("Papers")
            if papers and "Paper Name" in papers["headers"] and "Source File" in papers["headers"]:
                ni = papers["headers"].index("Paper Name")
                fi = papers["headers"].index("Source File")
                for row in papers["rows"]:
                    paper_files.setdefault(str(row[ni]), str(row[fi]))

            self.sheets, self.paper_files, self.mtime = sheets, paper_files, mtime

    def query(self, name: str, q: str, sort: int | None, desc: bool, page: int) -> dict:
        self.refresh()
        sheet = self.sheets.get(name)
        if sheet is None:
            return {"error": f"no sheet named {name!r}"}
        headers, rows = sheet["headers"], sheet["rows"]

        if q:
            terms = q.lower().split()
            rows = [r for r in rows if all(any(t in str(v).lower() for v in r) for t in terms)]

        if sort is not None and 0 <= sort < len(headers):
            def key(r):
                v = r[sort]
                # Numbers before text, numbers compared numerically
                return (0, v, "") if isinstance(v, (int, float)) else (1, 0, str(v).lower())
            rows = sorted(rows, key=key, reverse=desc)

        total = len(rows)
        pages = max((total + PAGE_SIZE - 1) // PAGE_SIZE, 1)
        page = min(max(page, 1), pages)
        chunk = rows[(page - 1) * PAGE_SIZE : page * PAGE_SIZE]

        paper_col = next((headers.index(c) for c in PAPER_COLUMNS if c in headers), None)
        if "Source File" in headers:
            src_col = headers.index("Source File")
            pdfs = [str(r[src_col]) or None for r in chunk]
        else:
            # Spreadsheets from older analyzer versions: look the file up by paper name
            pdfs = [self.paper_files.get(str(r[paper_col])) if paper_col is not None else None for r in chunk]
        return {
            "headers": headers,
            "rows": chunk,
            "pdfs": pdfs,
            "paper_col": paper_col,
            "total": total,
            "page": page,
            "pages": pages,
        }


def safe_file(base: Path, rel: str) -> Path | None:
    """Resolve rel inside base, refusing anything that escapes it."""
    path = (base / rel).resolve()
    if base.resolve() not in path.parents or not path.is_file():
        return None
    return path


def make_handler(db: Database, pdf_dir: Path, figures_dir: Path):
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            pass

        def send_body(self, status: int, body: bytes, content_type: str):
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def send_json(self, data, status=200):
            self.send_body(status, json.dumps(data, default=str).encode(), "application/json")

        def send_file(self, path: Path | None, download: bool = False):
            if path is None:
                self.send_body(404, b"Not found", "text/plain")
                return
            ctype = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
            body = path.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            if download:
                self.send_header("Content-Disposition", f'attachment; filename="{path.name}"')
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            url = urlparse(self.path)
            params = {k: v[0] for k, v in parse_qs(url.query).items()}
            try:
                if url.path == "/":
                    self.send_body(200, PAGE_HTML.encode(), "text/html; charset=utf-8")
                elif url.path == "/api/sheets":
                    db.refresh()
                    self.send_json([{"name": n, "count": len(s["rows"])} for n, s in db.sheets.items()])
                elif url.path == "/api/sheet":
                    sort = int(params["sort"]) if params.get("sort", "").isdigit() else None
                    result = db.query(
                        params.get("name", ""), params.get("q", "").strip(), sort,
                        params.get("desc") == "1", int(params.get("page", "1") or 1),
                    )
                    # No link for papers whose PDF isn't in --pdf-dir
                    if "pdfs" in result:
                        result["pdfs"] = [f if f and safe_file(pdf_dir, f) else None for f in result["pdfs"]]
                    self.send_json(result)
                elif url.path == "/download":
                    self.send_file(db.xlsx_path, download=True)
                elif url.path.startswith("/pdf/"):
                    self.send_file(safe_file(pdf_dir, unquote(url.path[len("/pdf/"):])))
                elif url.path.startswith("/figures/"):
                    self.send_file(safe_file(figures_dir, unquote(url.path[len("/figures/"):])))
                else:
                    self.send_body(404, b"Not found", "text/plain")
            except (BrokenPipeError, ConnectionResetError):
                pass
            except Exception as e:
                self.send_json({"error": str(e)}, status=500)

    return Handler


PAGE_HTML = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Radiation Paper Database</title>
<style>
  :root { --bg:#fff; --fg:#1d2330; --muted:#667085; --line:#e3e6ec; --head:#2f5496; --head-fg:#fff;
          --row:#f7f8fa; --accent:#2f5496; --hover:#eef2fb; }
  @media (prefers-color-scheme: dark) {
    :root { --bg:#14171d; --fg:#e4e7ec; --muted:#98a2b3; --line:#2a2f3a; --head:#22355e; --head-fg:#e4e7ec;
            --row:#1a1e25; --accent:#8fb0ff; --hover:#1f2633; }
  }
  * { box-sizing: border-box; }
  body { margin:0; font:14px/1.4 system-ui, sans-serif; background:var(--bg); color:var(--fg); }
  header { padding:16px; display:flex; flex-wrap:wrap; gap:12px; align-items:center; border-bottom:1px solid var(--line); }
  h1 { font-size:18px; margin:0 12px 0 0; }
  #download { margin-left:auto; white-space:nowrap; }
  #tabs { display:flex; flex-wrap:wrap; gap:4px; }
  #tabs button { border:1px solid var(--line); background:var(--bg); color:var(--fg); padding:6px 10px;
                 border-radius:6px; cursor:pointer; font:inherit; }
  #tabs button.active { background:var(--accent); border-color:var(--accent); color:var(--bg); }
  #tabs small { opacity:.7; margin-left:4px; }
  .bar { padding:12px 16px; display:flex; gap:12px; align-items:center; flex-wrap:wrap; }
  #q { flex:1; min-width:200px; max-width:480px; padding:8px 10px; border:1px solid var(--line); border-radius:6px;
       background:var(--bg); color:var(--fg); font:inherit; }
  #status { color:var(--muted); }
  .pager button { font:inherit; padding:4px 10px; border:1px solid var(--line); background:var(--bg); color:var(--fg);
                  border-radius:6px; cursor:pointer; }
  .pager button:disabled { opacity:.4; cursor:default; }
  .wrap { overflow:auto; max-height:calc(100vh - 140px); margin:0 16px 16px; border:1px solid var(--line); border-radius:6px; }
  table { border-collapse:collapse; width:max-content; min-width:100%; }
  th { position:sticky; top:0; background:var(--head); color:var(--head-fg); text-align:left; padding:8px;
       cursor:pointer; white-space:nowrap; user-select:none; }
  td { padding:6px 8px; border-top:1px solid var(--line); vertical-align:top; max-width:420px; }
  tr:nth-child(even) td { background:var(--row); }
  tr:hover td { background:var(--hover); }
  a { color:var(--accent); }
</style>
</head>
<body>
<header><h1>Radiation Paper Database</h1><div id="tabs"></div>
  <a id="download" href="download">Download spreadsheet (.xlsx)</a></header>
<div class="bar">
  <input id="q" type="search" placeholder="Search this sheet (all words must match)">
  <span id="status"></span>
  <span class="pager"><button id="prev">&larr; Prev</button> <button id="next">Next &rarr;</button></span>
</div>
<div class="wrap"><table><thead id="thead"></thead><tbody id="tbody"></tbody></table></div>
<script>
const state = { sheet: null, q: "", sort: null, desc: false, page: 1, pages: 1 };
const $ = id => document.getElementById(id);

function esc(s) { return String(s).replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c])); }

async function loadTabs() {
  const sheets = await (await fetch("api/sheets")).json();
  $("tabs").innerHTML = "";
  for (const s of sheets) {
    const b = document.createElement("button");
    b.innerHTML = esc(s.name) + "<small>" + s.count + "</small>";
    b.onclick = () => { state.sheet = s.name; state.sort = null; state.page = 1; render(); };
    b.dataset.name = s.name;
    $("tabs").appendChild(b);
  }
  if (!state.sheet && sheets.length) state.sheet = sheets[0].name;
  render();
}

async function render() {
  for (const b of $("tabs").children) b.classList.toggle("active", b.dataset.name === state.sheet);
  const p = new URLSearchParams({ name: state.sheet, q: state.q, page: state.page,
                                  sort: state.sort ?? "", desc: state.desc ? "1" : "0" });
  const d = await (await fetch("api/sheet?" + p)).json();
  if (d.error) { $("status").textContent = d.error; return; }
  state.page = d.page; state.pages = d.pages;

  const img = d.headers.indexOf("Image"), pageCol = d.headers.indexOf("Page");
  $("thead").innerHTML = "<tr>" + d.headers.map((h, i) =>
    `<th data-i="${i}">${esc(h)}${state.sort === i ? (state.desc ? " &#9660;" : " &#9650;") : ""}</th>`).join("") + "</tr>";
  $("tbody").innerHTML = d.rows.map((r, ri) => "<tr>" + r.map((v, i) => {
    let cell = esc(v);
    const pdf = d.pdfs[ri];
    if (i === d.paper_col && pdf) {
      const anchor = pageCol >= 0 && r[pageCol] ? "#page=" + r[pageCol] : "";
      cell = `<a href="pdf/${encodeURIComponent(pdf)}${anchor}" target="_blank">${cell}</a>`;
    } else if (i === img && v) {
      // Stored as <figures-dir>/<paper>/<file>.png — serve the last two parts
      const rel = String(v).split(/[\\\\/]/).slice(-2).map(encodeURIComponent).join("/");
      cell = `<a href="figures/${rel}" target="_blank">view image</a>`;
    }
    return "<td>" + cell + "</td>";
  }).join("") + "</tr>").join("");

  $("status").textContent = `${d.total.toLocaleString()} rows · page ${d.page} of ${d.pages}`;
  $("prev").disabled = d.page <= 1;
  $("next").disabled = d.page >= d.pages;
}

$("thead").onclick = e => {
  const th = e.target.closest("th"); if (!th) return;
  const i = +th.dataset.i;
  state.desc = state.sort === i ? !state.desc : false;
  state.sort = i; state.page = 1; render();
};
let timer;
$("q").oninput = () => { clearTimeout(timer); timer = setTimeout(() => { state.q = $("q").value; state.page = 1; render(); }, 250); };
$("prev").onclick = () => { state.page--; render(); };
$("next").onclick = () => { state.page++; render(); };
loadTabs();
</script>
</body>
</html>
"""


def main():
    parser = argparse.ArgumentParser(description="Serve the radiation paper spreadsheet as a web page.")
    parser.add_argument("--xlsx", default="rad_effects_database.xlsx",
                        help="Spreadsheet made by the analyzer (default: rad_effects_database.xlsx)")
    parser.add_argument("--pdf-dir", default="pdfs", help="Directory with the paper PDFs (default: pdfs)")
    parser.add_argument("--figures-dir", default="figures",
                        help="Directory with extracted figure images (default: figures)")
    parser.add_argument("--host", default="127.0.0.1",
                        help="Address to listen on; 0.0.0.0 shares it on your network (default: 127.0.0.1)")
    parser.add_argument("--port", type=int, default=8000, help="Port (default: 8000)")
    args = parser.parse_args()

    xlsx = Path(args.xlsx).expanduser()
    if not xlsx.is_file():
        raise SystemExit(f"Error: spreadsheet '{xlsx}' not found. Run rad_paper_analyzer_local.py first.")
    db = Database(xlsx)
    db.refresh()

    pdf_dir = Path(args.pdf_dir).expanduser()
    figures_dir = Path(args.figures_dir).expanduser()
    handler = make_handler(db, pdf_dir, figures_dir)
    server = ThreadingHTTPServer((args.host, args.port), handler)

    shown_host = "localhost" if args.host in ("127.0.0.1", "0.0.0.0") else args.host
    counts = ", ".join(f"{n}: {len(s['rows'])}" for n, s in db.sheets.items())
    print(f"Serving {xlsx} ({counts})")
    print(f"  PDFs:    {pdf_dir}/")
    print(f"  Figures: {figures_dir}/")
    print(f"  Open:    http://{shown_host}:{args.port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print()


if __name__ == "__main__":
    main()
