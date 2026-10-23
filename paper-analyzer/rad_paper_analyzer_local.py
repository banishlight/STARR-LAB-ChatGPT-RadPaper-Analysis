#!/usr/bin/env python3
"""
STARR-LAB Radiation Paper Analyzer — Local Model Edition
==========================================================
Same as rad_paper_analyzer.py but runs against a LOCAL LLM via Ollama
(or any OpenAI-compatible local server like vLLM, llama.cpp, LM Studio).

No API key needed. No data leaves your machine.

Requirements:
    pip install openai pymupdf openpyxl

Setup (Ollama — recommended):
    1. Install Ollama:          https://ollama.com/download
    2. Pull a model:            ollama pull qwen2.5:32b
    3. Place PDFs in ./pdfs/
    4. Run:                     python rad_paper_analyzer_local.py

The script talks to Ollama's OpenAI-compatible endpoint at localhost:11434.
For other servers (vLLM, llama.cpp, LM Studio), use --api-base to point
at their endpoint.
"""

import os
import sys
import base64
import json
import glob
import re
import time
import shutil
import argparse
from collections import Counter
from pathlib import Path

import pymupdf
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ─── Model recommendations for 24 GB VRAM (RTX 3090 / 4090) ───
MODEL_RECOMMENDATIONS = """
┌─────────────────────────────┬───────────┬─────────┬──────────────────────────┐
│ Model (ollama pull ...)     │ VRAM ~    │ Quality │ Notes                    │
├─────────────────────────────┼───────────┼─────────┼──────────────────────────┤
│ qwen2.5:32b                │ ~20 GB    │ ★★★★★  │ Best quality, tight fit   │
│ qwen2.5:14b                │ ~10 GB    │ ★★★★   │ Faster, more headroom     │
│ mistral-small:24b           │ ~14 GB    │ ★★★★   │ Strong all-rounder        │
│ llama3.1:8b                │ ~5 GB     │ ★★★    │ Fast, but less accurate   │
│ qwen3:30b-a3b              │ ~19 GB    │ ★★★★   │ MoE — fast + capable      │
│ gemma3:27b                 │ ~17 GB    │ ★★★★   │ Good at extraction        │
└─────────────────────────────┴───────────┴─────────┴──────────────────────────┘

For best results with radiation paper extraction, qwen2.5:32b or qwen2.5:14b
are recommended. They handle structured JSON output and technical content well.

Figure digitizing (--figures) needs a vision model, set with --vision-model:
  qwen2.5vl:7b   ~6 GB   good at reading charts (default)
  gemma3:27b     ~17 GB  slower, also handles images
"""


# ─── Prompts (tuned for local models — shorter, more explicit) ───

SYSTEM_PROMPT = """You are a radiation-effects specialist who extracts data from research papers.
You ALWAYS respond with a single valid JSON object. No other text, no markdown, no explanation."""

EXTRACTION_PROMPT = """Extract radiation test data from this paper into JSON.

Return this exact JSON structure (use null for missing values):

{
  "paper_name": "string",
  "year": 0,
  "authors": ["string"],
  "objective": "string",
  "parts": [
    {
      "name": "part number or device name",
      "type": "component type (FPGA, ADC, memory, etc.)",
      "manufacturer": "string",
      "other_details": "string or null",
      "tests": [
        {
          "test_type": "TID or SEE or DD",
          "source": "radiation source or null",
          "max_fluence": "string or null",
          "energy_levels": "string or null",
          "facility_name": "string or null",
          "environmental_conditions": "string or null",
          "terrestrial": true,
          "flight": false,
          "max_tid": "string or null",
          "dose_rate": "string or null",
          "eldrs": null,
          "dose_to_failure": "string or null",
          "increased_power_usage": null,
          "power_usage_description": "string or null",
          "failing_time": "string or null",
          "see_type": "ONE code from the SEE type list below, or null",
          "amplitude": "string or null",
          "duration": "string or null",
          "cross_section_saturation": "string or null",
          "cross_section_threshold": "string or null",
          "cross_section_type": "string or null",
          "damage_level": "string or null",
          "damage_description": "string or null",
          "parametrics": [
            {
              "parameter": "ONE code from the parametric list below",
              "label": "parameter name as written in the paper",
              "pre": "value before irradiation, with unit, or null",
              "post": "value after irradiation, with unit, or null",
              "at_dose": "dose or fluence of the post value, or null",
              "spec_limit": "datasheet/spec limit, or null",
              "notes": "string or null"
            }
          ],
          "special_notes": "string or null"
        }
      ]
    }
  ]
}

SEE type list (use exactly these codes for "see_type"):
  SEU   - Single Event Upset (single bit flip)
  MBU   - Multiple Bit Upset (several bits in one word)
  MCU   - Multiple Cell Upset (several cells, possibly different words)
  SET   - Single Event Transient (voltage glitch; includes ASET/DSET)
  SEFI  - Single Event Functional Interrupt (hang, reset, mode change)
  SEL   - Single Event Latchup (high current, needs power cycle)
  HCE   - High Current Event / micro-latchup
  SESB  - Single Event Stuck Bit / Single Hard Error (SHE, SEHE)
  SEB   - Single Event Burnout
  SEGR  - Single Event Gate Rupture
  SEDR  - Single Event Dielectric Rupture
  SES   - Single Event Snapback
  OTHER - an SEE not in this list (describe it in "special_notes")

Parametric list (use exactly these codes for "parameter"):
  IBIAS   - input bias current (op-amps, comparators, bipolar parts)
  BETA    - transistor current gain (beta, hFE)
  INL     - integral nonlinearity (ADCs, DACs)
  DNL     - differential nonlinearity (ADCs, DACs)
  ISUPPLY - supply current (Icc, Idd, quiescent current)
  VOUT    - output voltage (regulators, references, VOH/VOL)
  OTHER   - any other measured electrical parameter shift (offset voltage,
            leakage current, ...); name it in "label"

Rules:
- One entry per component in "parts". One entry per test in "tests".
- For SEE tests: make a SEPARATE test entry for EACH SEE type the paper
  reports for that component (e.g. a part with SEU and SEL results gets two
  SEE entries). Only list SEE types the paper actually observed or tested for.
- "parametrics": one entry per electrical parameter the paper reports a
  value or shift for under irradiation. If several doses are reported, give
  the highest dose as "post" and mention the others in "notes". Use [] if
  the paper reports none.
- Use null for values not found in the paper.
- Return ONLY the JSON object. No other text.

Paper text:

"""

FIGURE_PROMPT = """This image is a figure from a radiation-effects research paper.

Caption: {caption}
Components studied in the paper: {parts}

Return this exact JSON structure (use null for missing values):

{{
  "kind": "data_plot or diagram or photo or table or other",
  "title": "string or null",
  "x_axis": {{"label": "string", "unit": "string or null", "scale": "linear or log"}},
  "y_axis": {{"label": "string", "unit": "string or null", "scale": "linear or log"}},
  "series": [
    {{"label": "legend entry or description", "points": [[0.0, 0.0]]}}
  ],
  "related_part": "which component above the figure is about, or null",
  "notes": "string or null"
}}

Rules:
- Only fill "series" when kind is "data_plot". Otherwise use an empty list.
- Read values from the axis tick labels. On a log axis, interpolate in log space.
- Write numbers as plain numbers, e.g. 2e-3 not "2.E-03".
- One series per legend entry. For markers, give every marker. For lines,
  give 10-20 points including the endpoints and any bends.
- For multi-panel figures, prefix each series label with its panel, e.g. "(a) SEU".
- Return ONLY the JSON object. No other text.
"""

# Bump when extract_figures changes so saved papers get their figures re-cropped
FIGURES_VERSION = 2

# Matches captions like "Figure 3.", "Fig. 3:", "FIGURE 3 -"
CAPTION_RE = re.compile(r"^\s*(?:fig\.?|figure)\s*(\d+)\s*[.:\-–—]?\s*", re.I)

REFERENCES_RE = re.compile(
    r"^\s*(?:[0-9IVX]+\.?\s*)?(?:references|bibliography|reference list)\s*$",
    re.I | re.M,
)


def extract_text_from_pdf(pdf_path: str) -> str:
    """Extract text from a PDF using PyMuPDF."""
    doc = pymupdf.open(pdf_path)
    text_parts = []
    for page in doc:
        text_parts.append(page.get_text())
    doc.close()
    return "\n".join(text_parts)


def strip_references(text: str) -> str:
    """Drop the reference list so the context budget goes to the paper body."""
    matches = [m for m in REFERENCES_RE.finditer(text) if m.start() > len(text) * 0.4]
    if not matches:
        return text
    return text[: matches[-1].start()]


def _image_key(info: dict):
    # Same embedded image (xref) or same placement on the page
    return ("img", info["xref"]) if info["xref"] else ("img", tuple(round(v) for v in info["bbox"]))


def _drawing_key(d: dict):
    return ("draw", tuple(round(v) for v in d["rect"]))


def _graphic_regions(page, images: list[dict], drawings: list[dict], template: set) -> list:
    """Bounding boxes of images and vector drawings big enough to be figures."""
    page_area = page.rect.get_area()
    # Skip template decoration (slide sidebars, logos, header bars) that repeats across pages
    rects = [pymupdf.Rect(i["bbox"]) for i in images if _image_key(i) not in template]
    # Ignore page-sized background shapes (common in slide decks) before clustering
    drawings = [d for d in drawings
                if d["rect"].get_area() < 0.5 * page_area and _drawing_key(d) not in template]
    if drawings:
        rects += [pymupdf.Rect(r) for r in page.cluster_drawings(drawings=drawings)]
    rects = [r & page.rect for r in rects]
    rects = [r for r in rects if r.width > 40 and r.height > 40 and r.get_area() > 0.02 * page_area]
    # Long thin strips are banners/title art/sidebars, not figures
    rects = [r for r in rects if max(r.width / r.height, r.height / r.width) < 6]

    # Merge regions that touch or nearly touch (multi-panel figures)
    merged = True
    while merged:
        merged = False
        for i in range(len(rects)):
            for j in range(i + 1, len(rects)):
                if (rects[i] + (-25, -25, 25, 25)).intersects(rects[j]):
                    rects[i] |= rects.pop(j)
                    merged = True
                    break
            if merged:
                break
    return rects


def extract_figures(pdf_path: str, out_dir: Path, dpi: int = 200) -> list[dict]:
    """Find figures by pairing graphic regions with their captions, save as PNGs."""
    doc = pymupdf.open(pdf_path)
    page_graphics = [(page.get_image_info(xrefs=True), page.get_drawings()) for page in doc]

    # Anything appearing on 3+ pages and at least 30% of pages is page template
    counts = Counter()
    for images, drawings in page_graphics:
        counts.update({_image_key(i) for i in images} | {_drawing_key(d) for d in drawings})
    template = {k for k, c in counts.items() if c >= 3 and c >= 0.3 * len(doc)}

    found = []
    for pno, page in enumerate(doc, 1):
        regions = _graphic_regions(page, *page_graphics[pno - 1], template)
        if not regions:
            continue
        caps = []
        for b in page.get_text("blocks"):
            if b[6] == 0 and (m := CAPTION_RE.match(b[4])):
                caps.append((pymupdf.Rect(b[:4]), m.group(1), " ".join(b[4].split())[:500]))

        # A caption inside a region means the region swallowed the caption (and
        # maybe content below it); the figure is the part above the caption.
        split = []
        while regions:
            r = regions.pop()
            inside = [c for c in caps if r.contains((c[0].tl + c[0].br) / 2)]
            if not inside:
                split.append(r)
                continue
            top = min(inside, key=lambda c: c[0].y0)[0]
            above = pymupdf.Rect(r.x0, r.y0, r.x1, top.y0)
            below = pymupdf.Rect(r.x0, top.y1, r.x1, r.y1)
            if above.height > 40:
                split.append(above)
            if below.height > 40:
                regions.append(below)

        # Assign each region to the nearest caption, preferring captions below
        groups = {}
        for r in split:
            best, best_d = None, page.rect.height / 3
            for idx, (crect, _, _) in enumerate(caps):
                if crect.y0 >= r.y1 - 20:
                    d = max(crect.y0 - r.y1, 0)
                elif crect.y1 <= r.y0 + 20:
                    d = (r.y0 - crect.y1) * 1.5
                else:
                    continue
                if d < best_d:
                    best, best_d = idx, d
            groups.setdefault(best, []).append(r)

        for idx, rs in groups.items():
            if idx is None:
                # Uncaptioned graphics (e.g. plots on slides) — keep only large ones
                for r in rs:
                    if r.get_area() > 0.08 * page.rect.get_area():
                        found.append((pno, None, "", r))
                continue
            rect = rs[0]
            for r in rs[1:]:
                rect |= r
            found.append((pno, caps[idx][1], caps[idx][2], rect))

    figures = []
    if found:
        out_dir.mkdir(parents=True, exist_ok=True)
    seen_ids = set()
    for k, (pno, num, caption, rect) in enumerate(found, 1):
        page = doc[pno - 1]
        clip = (rect + (-18, -18, 18, 18)) & page.rect   # padding keeps axis labels
        fig_id = f"fig{num}" if num else f"p{pno}_img{k}"
        if fig_id in seen_ids:
            fig_id = f"{fig_id}_{k}"
        seen_ids.add(fig_id)
        # Cap the long side at ~1600 px so vision models don't choke on huge images
        scale = min(dpi, 1600 * 72 / max(clip.width, clip.height))
        path = out_dir / f"{fig_id}_p{pno}.png"
        page.get_pixmap(clip=clip, dpi=int(scale)).save(path)
        figures.append({"figure_id": fig_id, "page": pno, "caption": caption, "image_path": str(path)})
    doc.close()
    return figures


def extract_json_from_response(text: str) -> dict | None:
    """Try multiple strategies to pull valid JSON from a model response."""
    if not text:
        return None
    text = text.strip()

    # Strategy 1: direct parse
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    # Strategy 2: strip markdown fences
    cleaned = re.sub(r"^```(?:json)?\s*", "", text)
    cleaned = re.sub(r"\s*```\s*$", "", cleaned)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        pass

    # Strategy 3: find the first { ... } block (greedy)
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if match:
        try:
            return json.loads(match.group())
        except json.JSONDecodeError:
            pass

    # Strategy 4: find the outermost braces more carefully
    depth = 0
    start = None
    for i, ch in enumerate(text):
        if ch == "{":
            if depth == 0:
                start = i
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and start is not None:
                try:
                    return json.loads(text[start : i + 1])
                except json.JSONDecodeError:
                    start = None

    return None


def analyze_paper(
    client: OpenAI,
    text: str,
    model: str,
    max_chars: int,
    max_retries: int = 3,
    timeout: float = 300,
) -> dict | None:
    """Send paper text to local model and get structured JSON back with retries."""
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[Paper truncated due to context length limits]"

    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": EXTRACTION_PROMPT + text},
    ]
    return request_json(client, model, messages, "parts", max_retries, timeout, max_tokens=8192)


def digitize_figure(
    client: OpenAI,
    figure: dict,
    part_names: list[str],
    model: str,
    max_retries: int = 3,
    timeout: float = 300,
) -> dict | None:
    """Ask a vision model to classify a figure and read the data points off it."""
    with open(figure["image_path"], "rb") as f:
        image_b64 = base64.b64encode(f.read()).decode()
    prompt = FIGURE_PROMPT.format(
        caption=figure["caption"] or "(none)",
        parts=", ".join(part_names) or "(unknown)",
    )
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": [
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image_b64}"}},
        ]},
    ]
    data = request_json(client, model, messages, "kind", max_retries, timeout, max_tokens=4096)
    if data is None:
        return None

    # Keep only points the spreadsheet/database can use as numbers
    clean_series = []
    for s in data.get("series") or []:
        if not isinstance(s, dict):
            continue
        points = []
        for p in s.get("points") or []:
            try:
                points.append([float(p[0]), float(p[1])])
            except (TypeError, ValueError, IndexError):
                continue
        clean_series.append({"label": safe(s.get("label")), "points": points})
    data["series"] = clean_series
    return data


def request_json(
    client: OpenAI,
    model: str,
    messages: list[dict],
    required_key: str,
    max_retries: int,
    timeout: float,
    max_tokens: int,
) -> dict | None:
    """Call the model until it returns a JSON object containing required_key."""
    for attempt in range(1, max_retries + 1):
        try:
            response = client.chat.completions.create(
                model=model,
                temperature=0.1,
                max_tokens=max_tokens,
                messages=messages,
                timeout=timeout,
            )
            content = response.choices[0].message.content
            data = extract_json_from_response(content)
            if isinstance(data, dict) and required_key in data:
                return data
            if attempt < max_retries:
                print(f"    ⚠ Attempt {attempt}: got response but JSON was invalid, retrying...")
        except Exception as e:
            err_str = str(e)
            # Fatal errors — no point retrying
            if "not found" in err_str or "404" in err_str:
                print(f"    ✗ Model not found: {model}")
                print(f"      Pull it first: ollama pull {model}")
                sys.exit(1)
            if attempt < max_retries:
                print(f"    ⚠ Attempt {attempt} error: {e}, retrying...")
            else:
                print(f"    ✗ All {max_retries} attempts failed. Last error: {e}")

    return None


SEE_CODES = {"SEU", "MBU", "MCU", "SET", "SEFI", "SEL", "HCE", "SESB",
             "SEB", "SEGR", "SEDR", "SES", "OTHER"}

SEE_ALIASES = {
    "SBU": "SEU", "UPSET": "SEU", "SINGLE EVENT UPSET": "SEU", "BIT FLIP": "SEU",
    "MULTIPLE BIT UPSET": "MBU", "MULTIPLE CELL UPSET": "MCU",
    "ASET": "SET", "DSET": "SET", "TRANSIENT": "SET", "SINGLE EVENT TRANSIENT": "SET",
    "FUNCTIONAL INTERRUPT": "SEFI", "SINGLE EVENT FUNCTIONAL INTERRUPT": "SEFI",
    "LATCHUP": "SEL", "LATCH-UP": "SEL", "SINGLE EVENT LATCHUP": "SEL",
    "SINGLE EVENT LATCH-UP": "SEL", "MICRO-LATCHUP": "HCE", "HIGH CURRENT EVENT": "HCE",
    "SHE": "SESB", "SEHE": "SESB", "STUCK BIT": "SESB", "STUCKBIT": "SESB",
    "STUCK BITS": "SESB", "HARD ERROR": "SESB", "SINGLE HARD ERROR": "SESB",
    "BURNOUT": "SEB", "SINGLE EVENT BURNOUT": "SEB",
    "GATE RUPTURE": "SEGR", "SINGLE EVENT GATE RUPTURE": "SEGR",
    "DIELECTRIC RUPTURE": "SEDR", "SNAPBACK": "SES", "SESNAP": "SES",
}


def normalize_see_type(val) -> str:
    """Map whatever the model wrote onto one of the standard SEE codes."""
    if not val:
        return ""
    raw = str(val).strip().upper()
    if raw in SEE_CODES:
        return raw
    if raw in SEE_ALIASES:
        return SEE_ALIASES[raw]
    # Model sometimes returns lists like "SEU, SEL" — normalize each piece
    pieces = [p.strip() for p in re.split(r"[,/;]| and ", raw) if p.strip()]
    mapped = [p if p in SEE_CODES else SEE_ALIASES.get(p, p) for p in pieces]
    return ", ".join(dict.fromkeys(mapped))


PARAMETRIC_CODES = {"IBIAS", "BETA", "INL", "DNL", "ISUPPLY", "VOUT", "OTHER"}

PARAMETRIC_ALIASES = {
    "IB": "IBIAS", "I_BIAS": "IBIAS", "BIAS CURRENT": "IBIAS", "INPUT BIAS CURRENT": "IBIAS",
    "HFE": "BETA", "GAIN": "BETA", "CURRENT GAIN": "BETA", "Β": "BETA",
    "INTEGRAL NONLINEARITY": "INL", "INTEGRAL NON-LINEARITY": "INL",
    "DIFFERENTIAL NONLINEARITY": "DNL", "DIFFERENTIAL NON-LINEARITY": "DNL",
    "ICC": "ISUPPLY", "IDD": "ISUPPLY", "IQ": "ISUPPLY", "SUPPLY CURRENT": "ISUPPLY",
    "POWER SUPPLY CURRENT": "ISUPPLY", "QUIESCENT CURRENT": "ISUPPLY",
    "VO": "VOUT", "V_OUT": "VOUT", "VOH": "VOUT", "VOL": "VOUT", "OUTPUT VOLTAGE": "VOUT",
}


def normalize_parametric(val) -> str:
    """Map whatever the model wrote onto one of the standard parametric codes."""
    raw = safe(val).strip().upper()
    if raw in PARAMETRIC_CODES:
        return raw
    return PARAMETRIC_ALIASES.get(raw, "OTHER" if raw else "")


def safe(val) -> str:
    """Convert a value to a display string, handling None/bool/lists."""
    if val is None:
        return ""
    if isinstance(val, bool):
        return "Yes" if val else "No"
    if isinstance(val, list):
        return ", ".join(str(v) for v in val)
    return str(val)


def build_spreadsheet(results: list[dict], output_path: str):
    """Write extracted data to a formatted Excel workbook."""
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2F5496")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")

    def write_sheet(ws, headers, rows):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        for r_idx, row in enumerate(rows, 2):
            for col, val in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=col, value=val)
                cell.border = thin_border
                cell.alignment = wrap

    # ── Papers ──
    ws_papers = wb.active
    ws_papers.title = "Papers"
    paper_rows = []
    for i, r in enumerate(results, 1):
        d = r["data"]
        paper_rows.append([
            i, safe(d.get("paper_name")), d.get("year", ""),
            safe(d.get("authors")), safe(d.get("objective")),
            len(d.get("parts", [])), r["filename"],
        ])
    write_sheet(ws_papers,
        ["#", "Paper Name", "Year", "Authors", "Objective", "# Parts", "Source File"],
        paper_rows)
    for col, w in zip("ABCDEFG", [5, 40, 8, 30, 50, 8, 25]):
        ws_papers.column_dimensions[col].width = w

    # ── Components ──
    ws_parts = wb.create_sheet("Components")
    part_rows = []
    for r in results:
        d = r["data"]
        for part in d.get("parts", []):
            part_rows.append([
                safe(d.get("paper_name")), d.get("year", ""),
                safe(part.get("name")), safe(part.get("type")),
                safe(part.get("manufacturer")), safe(part.get("other_details")), r["filename"],
            ])
    write_sheet(ws_parts,
        ["Paper Name", "Year", "Part Name", "Type", "Manufacturer", "Other Details", "Source File"],
        part_rows)
    for col, w in zip("ABCDEFG", [35, 8, 25, 20, 20, 30, 25]):
        ws_parts.column_dimensions[col].width = w

    # ── TID Data ──
    ws_tid = wb.create_sheet("TID Data")
    tid_rows = []
    for r in results:
        d = r["data"]
        for part in d.get("parts", []):
            for test in part.get("tests", []):
                if (test.get("test_type") or "").upper() != "TID":
                    continue
                tid_rows.append([
                    safe(d.get("paper_name")), d.get("year", ""),
                    safe(part.get("name")), safe(test.get("source")),
                    safe(test.get("max_fluence")), safe(test.get("energy_levels")),
                    safe(test.get("facility_name")), safe(test.get("environmental_conditions")),
                    safe(test.get("terrestrial")), safe(test.get("flight")),
                    safe(test.get("max_tid")), safe(test.get("dose_rate")),
                    safe(test.get("eldrs")), safe(test.get("dose_to_failure")),
                    safe(test.get("increased_power_usage")), safe(test.get("power_usage_description")),
                    safe(test.get("failing_time")), safe(test.get("special_notes")), r["filename"],
                ])
    write_sheet(ws_tid, [
        "Paper", "Year", "Part", "Source", "Max Fluence", "Energy",
        "Facility", "Environment", "Terrestrial", "Flight",
        "Max TID", "Dose Rate", "ELDRS", "Dose to Failure",
        "Increased Power", "Power Description", "Failing Time", "Notes", "Source File",
    ], tid_rows)

    # ── SEE Data ──
    ws_see = wb.create_sheet("SEE Data")
    see_rows = []
    for r in results:
        d = r["data"]
        for part in d.get("parts", []):
            for test in part.get("tests", []):
                if (test.get("test_type") or "").upper() != "SEE":
                    continue
                see_rows.append([
                    safe(d.get("paper_name")), d.get("year", ""),
                    safe(part.get("name")), safe(test.get("source")),
                    normalize_see_type(test.get("see_type")), safe(test.get("max_fluence")),
                    safe(test.get("energy_levels")),
                    safe(test.get("facility_name")), safe(test.get("environmental_conditions")),
                    safe(test.get("terrestrial")), safe(test.get("flight")),
                    safe(test.get("amplitude")), safe(test.get("duration")),
                    safe(test.get("cross_section_saturation")),
                    safe(test.get("cross_section_threshold")),
                    safe(test.get("cross_section_type")),
                    safe(test.get("special_notes")), r["filename"],
                ])
    write_sheet(ws_see, [
        "Paper", "Year", "Part", "Source", "SEE Type", "Max Fluence", "Energy",
        "Facility", "Environment", "Terrestrial", "Flight",
        "Amplitude", "Duration",
        "Cross-Section Sat.", "Cross-Section Thresh.", "Cross-Section Type", "Notes", "Source File",
    ], see_rows)

    # ── DD Data ──
    ws_dd = wb.create_sheet("DD Data")
    dd_rows = []
    for r in results:
        d = r["data"]
        for part in d.get("parts", []):
            for test in part.get("tests", []):
                if (test.get("test_type") or "").upper() != "DD":
                    continue
                dd_rows.append([
                    safe(d.get("paper_name")), d.get("year", ""),
                    safe(part.get("name")), safe(test.get("source")),
                    safe(test.get("max_fluence")), safe(test.get("energy_levels")),
                    safe(test.get("facility_name")), safe(test.get("environmental_conditions")),
                    safe(test.get("terrestrial")), safe(test.get("flight")),
                    safe(test.get("damage_level")), safe(test.get("damage_description")),
                    safe(test.get("special_notes")), r["filename"],
                ])
    write_sheet(ws_dd, [
        "Paper", "Year", "Part", "Source", "Max Fluence", "Energy",
        "Facility", "Environment", "Terrestrial", "Flight",
        "Damage Level", "Damage Description", "Notes", "Source File",
    ], dd_rows)

    # ── All Tests (flat) ──
    ws_all = wb.create_sheet("All Tests (Flat)")
    flat_rows = []
    for r in results:
        d = r["data"]
        for part in d.get("parts", []):
            for test in part.get("tests", []):
                flat_rows.append([
                    safe(d.get("paper_name")), d.get("year", ""),
                    safe(part.get("name")), safe(part.get("manufacturer")),
                    safe(test.get("test_type")), safe(test.get("source")),
                    safe(test.get("facility_name")), safe(test.get("max_fluence")),
                    safe(test.get("energy_levels")), safe(test.get("environmental_conditions")),
                    safe(test.get("terrestrial")), safe(test.get("flight")),
                    safe(test.get("special_notes")), r["filename"],
                ])
    write_sheet(ws_all, [
        "Paper", "Year", "Part", "Manufacturer", "Test Type", "Source",
        "Facility", "Max Fluence", "Energy", "Environment",
        "Terrestrial", "Flight", "Notes", "Source File",
    ], flat_rows)

    # ── Parametric Data ──
    ws_param = wb.create_sheet("Parametric Data")
    param_rows = []
    for r in results:
        d = r["data"]
        for part in d.get("parts", []):
            for test in part.get("tests", []):
                for p in test.get("parametrics") or []:
                    if not isinstance(p, dict):
                        continue
                    param_rows.append([
                        safe(d.get("paper_name")), d.get("year", ""),
                        safe(part.get("name")), safe(part.get("type")),
                        safe(test.get("test_type")),
                        normalize_parametric(p.get("parameter")), safe(p.get("label")),
                        safe(p.get("pre")), safe(p.get("post")), safe(p.get("at_dose")),
                        safe(p.get("spec_limit")), safe(p.get("notes")), r["filename"],
                    ])
    write_sheet(ws_param, [
        "Paper", "Year", "Part", "Part Type", "Test Type",
        "Parameter", "Label in Paper", "Pre-Rad", "Post-Rad", "At Dose/Fluence",
        "Spec Limit", "Notes", "Source File",
    ], param_rows)

    data_sheets = [ws_tid, ws_see, ws_dd, ws_all, ws_param]

    # ── Figures / Figure Data (only when figures were extracted) ──
    if any(r.get("figures") for r in results):
        ws_figs = wb.create_sheet("Figures")
        fig_rows, point_rows = [], []
        for r in results:
            d = r["data"]
            for fig in r.get("figures", []):
                a = fig.get("analysis") or {}
                x, y = a.get("x_axis") or {}, a.get("y_axis") or {}
                series = a.get("series", [])
                fig_rows.append([
                    safe(d.get("paper_name")), fig["figure_id"], fig["page"],
                    fig["caption"], safe(a.get("kind")), safe(a.get("related_part")),
                    safe(x.get("label")), safe(x.get("unit")), safe(x.get("scale")),
                    safe(y.get("label")), safe(y.get("unit")), safe(y.get("scale")),
                    len(series), sum(len(s["points"]) for s in series),
                    safe(a.get("notes")), fig["image_path"], r["filename"],
                ])
                for s in series:
                    for px, py in s["points"]:
                        point_rows.append([
                            safe(d.get("paper_name")), fig["figure_id"],
                            safe(a.get("related_part")), s["label"],
                            px, safe(x.get("label")), safe(x.get("unit")),
                            py, safe(y.get("label")), safe(y.get("unit")), r["filename"],
                        ])
        write_sheet(ws_figs, [
            "Paper", "Figure", "Page", "Caption", "Kind", "Related Part",
            "X Label", "X Unit", "X Scale", "Y Label", "Y Unit", "Y Scale",
            "# Series", "# Points", "Notes", "Image", "Source File",
        ], fig_rows)
        # Make the image path clickable
        for row in ws_figs.iter_rows(min_row=2, min_col=16, max_col=16):
            row[0].hyperlink = row[0].value

        ws_points = wb.create_sheet("Figure Data")
        write_sheet(ws_points, [
            "Paper", "Figure", "Related Part", "Series",
            "X", "X Label", "X Unit", "Y", "Y Label", "Y Unit", "Source File",
        ], point_rows)
        data_sheets += [ws_figs, ws_points]

    # Set default column widths on data sheets
    for ws in data_sheets:
        for col_letter in [chr(c) for c in range(ord("A"), ord("A") + ws.max_column)]:
            ws.column_dimensions[col_letter].width = 16

    wb.save(output_path)


def check_ollama_running(base_url: str) -> bool:
    """Quick check if the server is reachable."""
    import urllib.request
    import urllib.error
    try:
        # Strip /v1 for the health check
        health_url = base_url.replace("/v1", "")
        req = urllib.request.Request(health_url, method="GET")
        urllib.request.urlopen(req, timeout=5)
        return True
    except Exception:
        return False


def save_progress(progress_dir: Path, entry: dict):
    safe_name = entry["filename"].replace(".pdf", ".json")
    with open(progress_dir / safe_name, "w") as f:
        json.dump(entry, f, indent=2)


def process_figures(client: OpenAI, results: list[dict], pdf_files: list[str], args, progress_dir: Path):
    """Second pass: extract figures and digitize plots for every analyzed paper.

    Runs after all text extraction so Ollama only swaps to the vision model once.
    Progress is saved per paper, so interrupted runs resume where they left off.
    """
    pdf_by_name = {os.path.basename(p): p for p in pdf_files}
    todo = [r for r in results if r["filename"] in pdf_by_name]
    print()
    print(f"Figure extraction: {len(todo)} paper(s) → {args.figures_dir}/")

    # Step 1: crop figures for every paper (fast) so they can be reviewed early.
    # Papers cropped by an older version of extract_figures are redone.
    for i, entry in enumerate(todo, 1):
        filename = entry["filename"]
        if entry.get("figures_version") == FIGURES_VERSION:
            continue
        out_dir = Path(args.figures_dir) / Path(filename).stem
        shutil.rmtree(out_dir, ignore_errors=True)
        try:
            entry["figures"] = extract_figures(pdf_by_name[filename], out_dir)
        except Exception as e:
            print(f"[{i}/{len(todo)}] {filename}: ✗ figure extraction failed: {e}")
            continue
        entry["figures_version"] = FIGURES_VERSION
        save_progress(progress_dir, entry)
        print(f"[{i}/{len(todo)}] {filename}: {len(entry['figures'])} figure(s)")

    if args.no_digitize:
        return

    # Step 2: digitize with the vision model
    todo = [r for r in todo if any("analysis" not in f for f in r.get("figures", []))]
    print()
    print(f"Digitizing figures with {args.vision_model}: {len(todo)} paper(s)")
    for i, entry in enumerate(todo, 1):
        pending = [f for f in entry["figures"] if "analysis" not in f]
        print(f"[{i}/{len(todo)}] {entry['filename']}: {len(pending)} figure(s)")
        part_names = [safe(p.get("name")) for p in entry["data"].get("parts", [])]
        for fig in pending:
            t0 = time.time()
            analysis = digitize_figure(
                client, fig, part_names, args.vision_model,
                max_retries=args.retries, timeout=args.timeout,
            )
            if analysis is None:
                print(f"    ✗ {fig['figure_id']}: could not read figure")
                continue
            fig["analysis"] = analysis
            n_points = sum(len(s["points"]) for s in analysis["series"])
            print(f"    ✓ {fig['figure_id']} ({analysis.get('kind')}): "
                  f"{len(analysis['series'])} series, {n_points} points in {time.time() - t0:.0f}s")
            # Save after each figure — vision calls are slow
            save_progress(progress_dir, entry)


def main():
    parser = argparse.ArgumentParser(
        description="Extract radiation effects data from papers using a LOCAL LLM.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=MODEL_RECOMMENDATIONS,
    )
    parser.add_argument(
        "--pdf-dir", default="pdfs",
        help="Directory containing PDF papers (default: pdfs)",
    )
    parser.add_argument(
        "--output", default="rad_effects_database.xlsx",
        help="Output Excel file (default: rad_effects_database.xlsx)",
    )
    parser.add_argument(
        "--model", default="qwen2.5:32b",
        help="Model name as known by your server (default: qwen2.5:32b)",
    )
    parser.add_argument(
        "--api-base", default="http://localhost:11434/v1",
        help="OpenAI-compatible API base URL (default: http://localhost:11434/v1 for Ollama)",
    )
    parser.add_argument(
        "--max-chars", type=int, default=30000,
        help="Max characters of paper text to send (default: 30000, ~8k tokens)",
    )
    parser.add_argument(
        "--retries", type=int, default=3,
        help="Retries per paper on JSON parse failure (default: 3)",
    )
    parser.add_argument(
        "--timeout", type=float, default=600,
        help="Timeout in seconds per API call (default: 600 — local models are slow)",
    )
    parser.add_argument(
        "--dump-json", action="store_true",
        help="Also save raw JSON responses per paper",
    )
    parser.add_argument(
        "--progress-dir", default=".progress",
        help="Directory for saving progress (default: .progress). Allows resuming interrupted runs.",
    )
    parser.add_argument(
        "--no-resume", action="store_true",
        help="Start fresh, ignoring any saved progress",
    )
    parser.add_argument(
        "--figures", action="store_true",
        help="Also extract figures from each paper and digitize plots with --vision-model",
    )
    parser.add_argument(
        "--vision-model", default="qwen2.5vl:7b",
        help="Vision model used to read data off figures (default: qwen2.5vl:7b)",
    )
    parser.add_argument(
        "--no-digitize", action="store_true",
        help="With --figures, only save figure images; skip the vision model",
    )
    parser.add_argument(
        "--figures-dir", default="figures",
        help="Directory for extracted figure images (default: figures)",
    )
    parser.add_argument(
        "--show-models", action="store_true",
        help="Show recommended models for 24 GB VRAM and exit",
    )
    args = parser.parse_args()

    if args.show_models:
        print(MODEL_RECOMMENDATIONS)
        sys.exit(0)

    # Check server
    if not check_ollama_running(args.api_base):
        print(f"Error: Cannot reach server at {args.api_base}")
        if "11434" in args.api_base:
            print()
            print("  Ollama doesn't appear to be running. Start it with:")
            print("    ollama serve")
            print()
            print(f"  Then pull a model:")
            print(f"    ollama pull {args.model}")
        else:
            print(f"  Make sure your local LLM server is running at {args.api_base}")
        sys.exit(1)

    # Find PDFs
    pdf_dir = Path(args.pdf_dir)
    if not pdf_dir.is_dir():
        print(f"Error: PDF directory '{pdf_dir}' not found.")
        print(f"  mkdir {pdf_dir} && cp /path/to/papers/*.pdf {pdf_dir}/")
        sys.exit(1)

    pdf_files = sorted(glob.glob(str(pdf_dir / "*.pdf")))
    if not pdf_files:
        print(f"Error: No PDF files found in '{pdf_dir}/'")
        sys.exit(1)

    print(f"STARR-LAB Radiation Paper Analyzer (Local)")
    print(f"{'=' * 45}")
    print(f"  Server:    {args.api_base}")
    print(f"  Model:     {args.model}")
    print(f"  PDFs:      {len(pdf_files)} papers in {pdf_dir}/")
    print(f"  Max chars: {args.max_chars:,} (~{args.max_chars // 4:,} tokens)")
    print(f"  Retries:   {args.retries}")
    if args.figures:
        print(f"  Figures:   {args.figures_dir}/ "
              f"({'images only' if args.no_digitize else 'digitize with ' + args.vision_model})")
    print(f"  Output:    {args.output}")
    print()

    # Set up progress directory for resume support
    progress_dir = Path(args.progress_dir)
    if args.no_resume and progress_dir.exists():
        shutil.rmtree(progress_dir)
        print("  Cleared previous progress.\n")
    progress_dir.mkdir(exist_ok=True)

    # Load any previously completed results
    results = []
    completed = set()
    if not args.no_resume:
        for pf in sorted(progress_dir.glob("*.json")):
            try:
                with open(pf) as f:
                    saved = json.load(f)
                results.append(saved)
                completed.add(saved["filename"])
            except (json.JSONDecodeError, KeyError):
                pass
        if completed:
            print(f"  Resuming: {len(completed)} paper(s) already done, "
                  f"{len(pdf_files) - len(completed)} remaining\n")

    client = OpenAI(base_url=args.api_base, api_key="not-needed")

    for i, pdf_path in enumerate(pdf_files, 1):
        filename = os.path.basename(pdf_path)

        if filename in completed:
            continue

        print(f"[{i}/{len(pdf_files)}] Processing: {filename}")

        # Extract text
        print(f"    Extracting text from PDF...")
        try:
            text = extract_text_from_pdf(pdf_path)
        except Exception as e:
            print(f"    ✗ Failed to read PDF: {e}")
            continue

        if len(text.strip()) < 100:
            print(f"    ✗ PDF is scanned/image-only (very little text extracted)")
            continue

        print(f"    Extracted {len(text):,} chars ({len(text) // 4:,} est. tokens)")
        body = strip_references(text)
        if len(body) < len(text):
            print(f"    Dropped reference list ({len(text) - len(body):,} chars)")
            text = body
        truncated_chars = max(len(text) - args.max_chars, 0)
        if truncated_chars:
            print(f"    ⚠ Truncating to {args.max_chars:,} chars — "
                  f"last {truncated_chars:,} chars ({truncated_chars * 100 // len(text)}%) not sent")

        # Analyze
        print(f"    Sending to {args.model}...")
        t0 = time.time()
        data = analyze_paper(
            client, text, args.model,
            max_chars=args.max_chars,
            max_retries=args.retries,
            timeout=args.timeout,
        )
        elapsed = time.time() - t0

        if data is None:
            print(f"    ✗ Analysis failed after {args.retries} attempts ({elapsed:.0f}s)")
            continue

        parts_count = len(data.get("parts", []))
        tests_count = sum(len(p.get("tests", [])) for p in data.get("parts", []))
        print(f"    ✓ Done in {elapsed:.1f}s — {parts_count} component(s), {tests_count} test(s)")

        entry = {"filename": filename, "data": data, "truncated_chars": truncated_chars}
        results.append(entry)

        # Save progress immediately so we can resume
        save_progress(progress_dir, entry)

        if args.dump_json:
            json_path = pdf_path.replace(".pdf", "_extracted.json")
            with open(json_path, "w") as f:
                json.dump(data, f, indent=2)

    if args.figures:
        process_figures(client, results, pdf_files, args, progress_dir)

    print()
    if not results:
        print("No papers were successfully processed.")
        sys.exit(1)

    print(f"Building spreadsheet: {args.output}")
    build_spreadsheet(results, args.output)

    total_parts = sum(len(r["data"].get("parts", [])) for r in results)
    total_tests = sum(
        len(t.get("tests", []))
        for r in results
        for t in r["data"].get("parts", [])
    )
    print(f"  ✓ {len(results)} paper(s), {total_parts} component(s), {total_tests} test(s)")
    print(f"  ✓ Saved to {args.output}")
    print()
    print("Done!")


if __name__ == "__main__":
    main()